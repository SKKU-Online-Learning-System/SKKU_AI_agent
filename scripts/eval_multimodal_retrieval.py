"""Live Qwen check: scanned chart, office diagrams, semantic retrieval and isolation.

PYTHONPATH=apps/backend .venv-app/bin/python scripts/eval_multimodal_retrieval.py
Requires the configured model servers and LibreOffice; never used by unit tests.
"""
from __future__ import annotations

from dataclasses import asdict
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import time

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.core.config import Settings
from app.models import Base, Course, CourseMaterial, User, UserRole
from app.services.material_processing_service import MaterialProcessingService
from app.services.rag_service import RagService


def main():
    settings = Settings()
    rows = []
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with TemporaryDirectory(prefix="multimodal-eval-") as directory, Session(engine) as db:
        root = Path(directory)
        db.add(User(id="teacher", name="Evaluation", email="eval@example.test",
                    external_auth_id="eval", role=UserRole.professor))
        db.add_all([Course(id=c, name=c, semester="eval", professor_id="teacher")
                    for c in ["charts", "office", "other"]])
        db.commit()
        image = Image.new("RGB", (1200, 800), "white")
        draw = ImageDraw.Draw(image)
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 36)
        draw.text((50, 35), "Training experiment: learning rate and loss", font=font, fill="black")
        for i, (rate, loss) in enumerate([("0.01", "2.4"), ("0.10", "1.2"), ("1.00", "8.7")]):
            y = 170 + 180*i
            draw.text((40, y), "rate " + rate, font=font, fill="black")
            width = int(float(loss)*70)
            draw.rectangle((270, y, 270+width, y+70), fill="#2563eb")
            draw.text((290+width, y), "loss " + loss, font=font, fill="black")
        image.save(root / "scan.pdf", "PDF")  # No PDF text layer.
        png = BytesIO()
        image.save(png, "PNG")
        ppt = Presentation()
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        slide.shapes.add_picture(BytesIO(png.getvalue()), Inches(0), Inches(0), width=Inches(10))
        ppt.save(root / "chart.pptx")
        doc = Document()
        doc.add_heading("Training results", 0)
        doc.add_picture(BytesIO(png.getvalue()), width=Inches(6))
        doc.save(root / "chart.docx")
        (root / "unrelated.txt").write_text("기회비용은 포기한 대안 중 가장 큰 가치이다.")
        for name, course in [("scan.pdf", "charts"), ("chart.pptx", "office"),
                             ("chart.docx", "office"), ("unrelated.txt", "other")]:
            path = root / name
            material = CourseMaterial(id=name, course_id=course, uploaded_by="teacher",
                file_name=name, original_file_name=name, file_type=path.suffix[1:],
                file_size=path.stat().st_size, storage_path=str(path))
            db.add(material)
            db.commit()
            start = time.perf_counter()
            outcome = MaterialProcessingService(db, settings).process_material(name)
            rows.append({"index": name, "chunks": outcome.chunk_count,
                         "seconds": round(time.perf_counter()-start, 3)})
        for course in ["charts", "office", "other"]:
            start = time.perf_counter()
            outcome = RagService(db, settings).retrieve(
                course, "학습률이 0.10일 때 그림에 표시된 손실 값은 얼마야?", top_k=2,
            )
            rows.append({"course": course, "seconds": round(time.perf_counter()-start, 3),
                         "results": [asdict(r) for r in outcome.results]})
            if course != "other":
                assert outcome.results and all("1.2" in r.chunk_text for r in outcome.results)
                assert all(r.has_image and r.page_number == 1 for r in outcome.results)
            else:
                assert not outcome.results
        print(json.dumps(rows, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
