from pathlib import Path

import pytest
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter

from app.models import CourseMaterial, CourseMaterialStatus
from app.services.chunking_service import ParagraphChunkingService
from app.services.document_parser_service import (
    DocumentParsingError,
    FileDocumentParserService,
)


def material(path: Path, file_type: str) -> CourseMaterial:
    return CourseMaterial(
        id="material-1",
        course_id="course-1",
        uploaded_by="professor-1",
        file_name=path.name,
        original_file_name=f"lecture.{file_type}",
        file_type=file_type,
        file_size=path.stat().st_size if path.exists() else 0,
        storage_path=str(path),
        processing_status=CourseMaterialStatus.pending,
    )


def text_pdf(content: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({content}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, item in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode())
        output.extend(item)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010} 00000 n \n".encode())
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    return bytes(output)


@pytest.mark.parametrize("encoding", ["utf-8", "cp949"])
def test_txt_parser_preserves_korean_text(tmp_path: Path, encoding: str) -> None:
    path = tmp_path / "lecture.txt"
    path.write_bytes("경사하강법 강의".encode(encoding))

    parsed = FileDocumentParserService().parse(material(path, "txt"))

    assert parsed.material_id == "material-1"
    assert parsed.course_id == "course-1"
    assert parsed.title == "lecture.txt"
    assert parsed.full_text == "경사하강법 강의"
    assert parsed.pages[0].page_number is None
    assert parsed.pages[0].text == parsed.full_text


def test_pdf_parser_preserves_page_number(tmp_path: Path) -> None:
    path = tmp_path / "lecture.pdf"
    path.write_bytes(text_pdf("Gradient descent lecture"))

    parsed = FileDocumentParserService().parse(material(path, "pdf"))

    assert parsed.full_text == "Gradient descent lecture"
    assert [(page.page_number, page.text) for page in parsed.pages] == [
        (1, "Gradient descent lecture")
    ]


def test_long_pdf_text_splits_without_losing_page_number(tmp_path: Path) -> None:
    path = tmp_path / "long-lecture.pdf"
    path.write_bytes(text_pdf("A" * 2200))

    parsed = FileDocumentParserService().parse(material(path, "pdf"))
    chunks = ParagraphChunkingService().create_chunks(parsed)

    assert [chunk.char_count for chunk in chunks] == [1000, 1000, 500]
    assert all(chunk.page_number == 1 for chunk in chunks)


def test_docx_parser_extracts_paragraphs(tmp_path: Path) -> None:
    path = tmp_path / "lecture.docx"
    document = DocxDocument()
    document.add_paragraph("첫 번째 문단")
    document.add_paragraph("두 번째 문단")
    document.save(path)

    parsed = FileDocumentParserService().parse(material(path, "docx"))

    assert parsed.full_text == "첫 번째 문단\n\n두 번째 문단"
    assert parsed.pages[0].page_number is None


def test_pptx_parser_uses_slide_numbers(tmp_path: Path) -> None:
    path = tmp_path / "lecture.pptx"
    presentation = Presentation()
    for text in ("첫 슬라이드", "둘째 슬라이드"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = text
    presentation.save(path)

    parsed = FileDocumentParserService().parse(material(path, "pptx"))

    assert parsed.full_text == "첫 슬라이드\n\n둘째 슬라이드"
    assert [page.page_number for page in parsed.pages] == [1, 2]


def test_blank_pdf_reports_text_extraction_failure(tmp_path: Path) -> None:
    path = tmp_path / "scanned.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with path.open("wb") as destination:
        writer.write(destination)

    with pytest.raises(
        DocumentParsingError,
        match="DOCUMENT_TEXT_EMPTY: 텍스트를 추출할 수 없습니다",
    ):
        FileDocumentParserService().parse(material(path, "pdf"))


def test_missing_and_unsupported_files_have_stable_error_codes(tmp_path: Path) -> None:
    missing = tmp_path / "missing.txt"
    unsupported = tmp_path / "lecture.hwp"
    unsupported.write_bytes(b"hwp")

    with pytest.raises(DocumentParsingError, match="MATERIAL_FILE_NOT_FOUND"):
        FileDocumentParserService().parse(material(missing, "txt"))
    with pytest.raises(DocumentParsingError, match="DOCUMENT_UNSUPPORTED_TYPE"):
        FileDocumentParserService().parse(material(unsupported, "hwp"))
