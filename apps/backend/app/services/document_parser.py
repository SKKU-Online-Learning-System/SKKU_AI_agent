"""Text extraction for uploaded course materials.

Each supported extension is handled by a small reader so the parser stays
replaceable. Optional libraries are imported lazily and missing ones fail with a
clear message instead of crashing the whole processing pipeline.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

SUPPORTED_EXTENSIONS = frozenset({".txt", ".pdf", ".docx", ".pptx"})

# PDF and PPTX extraction regularly emits C0 control characters. PostgreSQL text
# columns reject NUL outright ("text fields cannot contain NUL (0x00) bytes"),
# which fails the INSERT of every chunk of a material at once, and the rest are
# never meaningful lecture content. Tab, newline and carriage return are kept.
CONTROL_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


class DocumentParseError(Exception):
    """Base error for the extraction step, carrying a stable error code."""

    code = "DOCUMENT_PARSE_FAILED"


class UnsupportedDocumentTypeError(DocumentParseError):
    code = "DOCUMENT_UNSUPPORTED_TYPE"


class MaterialFileNotFoundError(DocumentParseError):
    code = "MATERIAL_FILE_NOT_FOUND"


class EmptyDocumentError(DocumentParseError):
    code = "DOCUMENT_TEXT_NOT_FOUND"


class DocumentParserUnavailableError(DocumentParseError):
    code = "DOCUMENT_PARSER_UNAVAILABLE"


@dataclass(frozen=True)
class ParsedPage:
    page_number: Optional[int]
    text: str


@dataclass(frozen=True)
class ParsedDocument:
    material_id: str
    course_id: str
    title: str
    pages: list[ParsedPage]

    @property
    def full_text(self) -> str:
        return "\n\n".join(page.text for page in self.pages if page.text.strip())


def sanitize_text(text: str) -> str:
    """Strip control characters that cannot be stored or read as text."""
    return CONTROL_CHARACTERS.sub("", text)


def parse_material(
    *,
    material_id: str,
    course_id: str,
    title: str,
    storage_path: str | Path,
) -> ParsedDocument:
    """Read the stored file and return its text grouped by page."""

    path = Path(storage_path)
    if not path.is_file():
        raise MaterialFileNotFoundError("자료 파일을 찾을 수 없습니다.")

    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocumentTypeError(f"지원하지 않는 파일 형식입니다: {extension or '알 수 없음'}")

    if extension == ".txt":
        pages = _read_txt(path)
    elif extension == ".pdf":
        pages = _read_pdf(path)
    elif extension == ".docx":
        pages = _read_docx(path)
    else:
        pages = _read_pptx(path)

    # Every extension converges here, so this is the one place that has to
    # guarantee storable text for chunks, embeddings and cited excerpts.
    cleaned = (ParsedPage(page.page_number, sanitize_text(page.text)) for page in pages)
    non_empty = [page for page in cleaned if page.text.strip()]
    if not non_empty:
        raise EmptyDocumentError(
            "파일에서 텍스트를 추출할 수 없습니다. 스캔 이미지 PDF는 지원하지 않습니다."
        )

    return ParsedDocument(
        material_id=material_id,
        course_id=course_id,
        title=title,
        pages=non_empty,
    )


def _read_txt(path: Path) -> list[ParsedPage]:
    raw = path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "cp949"):
        try:
            return [ParsedPage(page_number=1, text=raw.decode(encoding))]
        except UnicodeDecodeError:
            continue

    return [ParsedPage(page_number=1, text=raw.decode("utf-8", errors="replace"))]


def _read_pdf(path: Path) -> list[ParsedPage]:
    try:
        from pypdf import PdfReader
    except ImportError as error:  # pragma: no cover - depends on the local install
        raise DocumentParserUnavailableError(
            "PDF 파서(pypdf)가 설치되지 않아 처리할 수 없습니다."
        ) from error

    try:
        reader = PdfReader(str(path))
        return [
            ParsedPage(page_number=index, text=page.extract_text() or "")
            for index, page in enumerate(reader.pages, start=1)
        ]
    except DocumentParseError:
        raise
    except Exception as error:
        raise DocumentParseError(f"PDF를 읽는 중 오류가 발생했습니다: {error}") from error


def _read_docx(path: Path) -> list[ParsedPage]:
    try:
        from docx import Document
    except ImportError as error:  # pragma: no cover - depends on the local install
        raise DocumentParserUnavailableError(
            "DOCX 파서(python-docx)가 설치되지 않아 처리할 수 없습니다."
        ) from error

    try:
        document = Document(str(path))
        text = "\n\n".join(
            paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
        )
        return [ParsedPage(page_number=None, text=text)]
    except DocumentParseError:
        raise
    except Exception as error:
        raise DocumentParseError(f"DOCX를 읽는 중 오류가 발생했습니다: {error}") from error


def _read_pptx(path: Path) -> list[ParsedPage]:
    try:
        from pptx import Presentation
    except ImportError as error:  # pragma: no cover - depends on the local install
        raise DocumentParserUnavailableError(
            "PPTX 파서(python-pptx)가 설치되지 않아 처리할 수 없습니다."
        ) from error

    try:
        presentation = Presentation(str(path))
        pages: list[ParsedPage] = []
        for index, slide in enumerate(presentation.slides, start=1):
            texts = [
                shape.text
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False) and shape.text.strip()
            ]
            pages.append(ParsedPage(page_number=index, text="\n".join(texts)))
        return pages
    except DocumentParseError:
        raise
    except Exception as error:
        raise DocumentParseError(f"PPTX를 읽는 중 오류가 발생했습니다: {error}") from error
