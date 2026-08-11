from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Optional, Protocol

from docx import Document as DocxDocument
from pptx import Presentation
from pypdf import PdfReader

from app.models import CourseMaterial

# PDF and PPTX extraction regularly emits C0 control characters. PostgreSQL text
# columns reject NUL outright ("text fields cannot contain NUL (0x00) bytes"),
# which fails the whole INSERT of a material's chunks, and the rest are never
# meaningful lecture content. Tab, newline and carriage return are kept.
CONTROL_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


class DocumentParsingError(Exception):
    pass


def sanitize_text(text: str) -> str:
    """Strip control characters that cannot be stored or read as text."""
    return CONTROL_CHARACTERS.sub("", text)


@dataclass(frozen=True)
class ParsedPage:
    page_number: Optional[int]
    text: str


@dataclass(frozen=True)
class ParsedDocument:
    material_id: str
    course_id: str
    title: str
    pages: tuple[ParsedPage, ...]
    full_text: str


class DocumentParserService(Protocol):
    def parse(self, material: CourseMaterial) -> ParsedDocument:
        raise NotImplementedError


class FileDocumentParserService:
    def parse(self, material: CourseMaterial) -> ParsedDocument:
        path = Path(material.storage_path)
        if not path.is_file():
            raise DocumentParsingError(
                "MATERIAL_FILE_NOT_FOUND: Material file not found"
            )

        file_type = material.file_type.lower().lstrip(".")
        if file_type == "txt":
            pages = self._parse_txt(path)
        elif file_type == "pdf":
            pages = self._parse_pdf(path)
        elif file_type == "docx":
            pages = self._parse_docx(path)
        elif file_type == "pptx":
            pages = self._parse_pptx(path)
        else:
            raise DocumentParsingError(
                f"DOCUMENT_UNSUPPORTED_TYPE: Unsupported document type: {file_type}"
            )

        # Every format converges here, so this is the one place that has to
        # guarantee storable text for chunks, embeddings and cited excerpts.
        pages = tuple(
            cleaned_page
            for page in pages
            if (cleaned_page := ParsedPage(page.page_number, sanitize_text(page.text).strip())).text
        )
        full_text = "\n\n".join(page.text for page in pages)
        if not full_text:
            raise DocumentParsingError(
                "DOCUMENT_TEXT_EMPTY: 텍스트를 추출할 수 없습니다"
            )
        return ParsedDocument(
            material_id=material.id,
            course_id=material.course_id,
            title=material.original_file_name,
            pages=pages,
            full_text=full_text,
        )

    @staticmethod
    def _parse_txt(path: Path) -> tuple[ParsedPage, ...]:
        content = path.read_bytes()
        for encoding in ("utf-8-sig", "cp949"):
            try:
                return (ParsedPage(page_number=None, text=content.decode(encoding)),)
            except UnicodeDecodeError:
                continue
        raise DocumentParsingError(
            "DOCUMENT_ENCODING_ERROR: TXT encoding must be UTF-8 or CP949"
        )

    @staticmethod
    def _parse_pdf(path: Path) -> tuple[ParsedPage, ...]:
        try:
            reader = PdfReader(path)
            return tuple(
                ParsedPage(page_number=index, text=page.extract_text() or "")
                for index, page in enumerate(reader.pages, start=1)
            )
        except Exception as exc:
            raise DocumentParsingError(
                "DOCUMENT_PARSE_FAILED: PDF parsing failed"
            ) from exc

    @staticmethod
    def _parse_docx(path: Path) -> tuple[ParsedPage, ...]:
        try:
            document = DocxDocument(str(path))
            text = "\n\n".join(paragraph.text for paragraph in document.paragraphs)
            return (ParsedPage(page_number=None, text=text),)
        except Exception as exc:
            raise DocumentParsingError(
                "DOCUMENT_PARSE_FAILED: DOCX parsing failed"
            ) from exc

    @staticmethod
    def _parse_pptx(path: Path) -> tuple[ParsedPage, ...]:
        try:
            presentation = Presentation(str(path))
            return tuple(
                ParsedPage(
                    page_number=index,
                    text="\n\n".join(
                        shape.text
                        for shape in slide.shapes
                        if hasattr(shape, "text") and shape.text.strip()
                    ),
                )
                for index, slide in enumerate(presentation.slides, start=1)
            )
        except Exception as exc:
            raise DocumentParsingError(
                "DOCUMENT_PARSE_FAILED: PPTX parsing failed"
            ) from exc
